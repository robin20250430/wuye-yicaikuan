from datetime import datetime, timezone
from io import BytesIO
from uuid import uuid4

from docx import Document as WordDocument
from fastapi import APIRouter, HTTPException, Query, Response
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas
from sqlalchemy import select

from ...db.database import DocumentORM, SessionLocal
from ...schemas.common import DocumentGenerateRequest, GeneratedDocument
from ...services.ai_service import generate_document_with_ai_fallback

router = APIRouter()

TITLE_MAP = {
    "notice": "物业费催缴通知书",
    "deadline_notice": "限期缴费通知书",
    "lawyer_letter": "物业费催缴律师函（草稿）",
    "litigation_notice": "起诉前告知书（草稿）",
}


def persist_document(payload: DocumentGenerateRequest, generated: GeneratedDocument) -> GeneratedDocument:
    with SessionLocal() as db:
        row = DocumentORM(
            id=generated.id,
            document_type=generated.document_type,
            title=generated.title,
            content=generated.content,
            disclaimer=generated.disclaimer,
            owner_name=payload.owner_name,
            property_address=payload.property_address,
            overdue_amount=payload.overdue_amount,
        )
        db.add(row)
        db.commit()
    return generated


@router.post("/generate", response_model=GeneratedDocument)
def create_document(payload: DocumentGenerateRequest) -> GeneratedDocument:
    generated = generate_document_with_ai_fallback(payload)
    return persist_document(payload, generated)


@router.get("/templates")
def get_templates() -> list[dict]:
    return [
        {"id": "notice", "title": TITLE_MAP["notice"], "category": "通知函", "description": "用于提醒业主按时缴纳物业费", "price": "免费"},
        {"id": "deadline_notice", "title": TITLE_MAP["deadline_notice"], "category": "通知函", "description": "用于明确缴费期限和后续措施", "price": "免费"},
        {"id": "lawyer_letter", "title": TITLE_MAP["lawyer_letter"], "category": "律师函", "description": "适合转由律师审阅后发送", "price": "29.9"},
        {"id": "litigation_notice", "title": TITLE_MAP["litigation_notice"], "category": "诉讼材料", "description": "用于诉前告知和风险提示", "price": "199"},
    ]


def get_document_or_404(document_id: str) -> DocumentORM:
    with SessionLocal() as db:
        row = db.scalar(select(DocumentORM).where(DocumentORM.id == document_id))
        if not row:
            raise HTTPException(status_code=404, detail="文书不存在")
        # Detach a plain value copy before closing the session.
        return DocumentORM(
            id=row.id, document_type=row.document_type, title=row.title,
            content=row.content, disclaimer=row.disclaimer,
            owner_name=row.owner_name, property_address=row.property_address,
            overdue_amount=row.overdue_amount, created_at=row.created_at,
        )


def build_docx(row: DocumentORM) -> bytes:
    document = WordDocument()
    document.add_heading(row.title, level=1)
    for paragraph in row.content.split("\n"):
        document.add_paragraph(paragraph)
    document.add_paragraph("")
    document.add_paragraph(f"提示：{row.disclaimer}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def build_pdf(row: DocumentORM) -> bytes:
    output = BytesIO()
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pdf = canvas.Canvas(output)
    pdf.setFont("STSong-Light", 16)
    pdf.drawCentredString(300, 800, row.title)
    pdf.setFont("STSong-Light", 10)
    y = 770
    for line in (row.content + "\n\n提示：" + row.disclaimer).splitlines():
        if y < 45:
            pdf.showPage()
            pdf.setFont("STSong-Light", 10)
            y = 800
        pdf.drawString(50, y, line[:85])
        y -= 18
    pdf.save()
    return output.getvalue()


@router.get("/{document_id}/download")
def download_document(document_id: str, format: str = Query("pdf", pattern="^(pdf|docx|txt)$")) -> Response:
    row = get_document_or_404(document_id)
    if format == "pdf":
        content, media_type, extension = build_pdf(row), "application/pdf", "pdf"
    elif format == "docx":
        content, media_type, extension = build_docx(row), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
    else:
        content, media_type, extension = (row.content + "\n\n提示：" + row.disclaimer).encode("utf-8"), "text/plain; charset=utf-8", "txt"
    return Response(content=content, media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{document_id}.{extension}"'})
